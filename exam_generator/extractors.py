from pathlib import Path
from io import BytesIO


def extract_text(file) -> str:
    name = getattr(file, "name", str(file)).lower()
    data = file.read() if hasattr(file, "read") else Path(file).read_bytes()

    if name.endswith(".pdf"):
        return _from_pdf(data)
    if name.endswith(".pptx"):
        return _from_pptx(data)
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    return f"[지원 안 되는 형식: {name}]"


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data))
    return "\n\n".join((p.extract_text() or "") for p in reader.pages)


def _from_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                out.append(shape.text)
    return "\n".join(out)


def estimate_tokens(text: str) -> int:
    """한글 글자당 약 0.7 토큰 (대략값)."""
    return int(len(text) * 0.7)
